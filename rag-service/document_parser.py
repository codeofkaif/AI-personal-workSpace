"""Structure-aware, page-aware parsers for the RAG ingestion pipeline.

Native parsers handle reliable digital-text extraction first. Unstructured is used for
Office/HTML layout-aware extraction and as a fallback; OCR is limited to empty PDF pages.
"""
from __future__ import annotations

import io
import logging
import os
import re
import shutil
from collections import OrderedDict
from typing import List, Tuple

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".pptx", ".html", ".htm", ".txt", ".md", ".csv", ".json"}


def _with_structure(category: str, text: str) -> str:
    value = (text or "").replace("\x00", "").strip()
    if not value:
        return ""
    if category.lower() in {"title", "sectionheader", "header"}:
        return f"## {value}"
    if category.lower() == "table":
        return f"TABLE\n{value}"
    return value


def _unstructured_pages(file_bytes: bytes, filename: str) -> List[Tuple[int, str]]:
    """Extract elements and preserve their categories and page number where available."""
    from unstructured.partition.auto import partition

    elements = partition(file=io.BytesIO(file_bytes), file_filename=filename)
    grouped: OrderedDict[int, List[str]] = OrderedDict()
    for element in elements:
        text = str(element or "").strip()
        if not text:
            continue
        metadata = getattr(element, "metadata", None)
        page = getattr(metadata, "page_number", None) if metadata is not None else None
        try:
            page_no = int(page) if page is not None else 0
        except (TypeError, ValueError):
            page_no = 0
        category = type(element).__name__
        # Some Unstructured versions expose category separately from the element class.
        category = str(getattr(element, "category", None) or category)
        shaped = _with_structure(category, text)
        grouped.setdefault(page_no, []).append(shaped)
    return [(page, "\n\n".join(parts).strip()) for page, parts in grouped.items() if parts]


def _pdf_pages(file_bytes: bytes, ocr_enabled: bool, ocr_lang: str) -> List[Tuple[int, str]]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(file_bytes))
    result: List[Tuple[int, str]] = []
    for index, page in enumerate(reader.pages):
        text = (page.extract_text() or "").replace("\x00", "").strip()
        if not text and ocr_enabled and shutil.which("tesseract"):
            try:
                import fitz
                import pytesseract
                from PIL import Image
                with fitz.open(stream=file_bytes, filetype="pdf") as document:
                    pixmap = document.load_page(index).get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
                    image = Image.frombytes("RGB", [pixmap.width, pixmap.height], pixmap.samples)
                    text = (pytesseract.image_to_string(
                        image, lang=ocr_lang, config="--psm 3"
                    ) or "").replace("\x00", "").strip()
            except Exception:
                logger.exception("PDF OCR failed on page %s of %s", index + 1, filename_for_log(file_bytes))
        result.append((index + 1, text))
    if any(text.strip() for _, text in result):
        return result
    # Optional layout-aware OCR fallback. Tesseract may not be installed or able to read a scan.
    try:
        fallback = _unstructured_pages(file_bytes, "fallback.pdf")
        if fallback:
            return fallback
    except Exception:
        logger.info("Unstructured PDF fallback was unavailable; no readable PDF text extracted.")
    return result


def filename_for_log(_file_bytes: bytes) -> str:
    # Do not log document content or derive identifiers from private file bytes.
    return "uploaded document"


def _docx_native(file_bytes: bytes) -> List[Tuple[int, str]]:
    from docx import Document

    document = Document(io.BytesIO(file_bytes))
    parts: List[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = str(getattr(paragraph.style, "name", "") or "").lower()
        if style.startswith("heading") or style in {"title", "subtitle"}:
            parts.append(f"## {text}")
        else:
            parts.append(text)
    for table in document.tables:
        rows = [
            " | ".join((cell.text or "").strip().replace("\n", " / ") for cell in row.cells)
            for row in table.rows
        ]
        if rows:
            parts.append("TABLE\n" + "\n".join(rows))
    content = "\n\n".join(parts).strip()
    return [(0, content)] if content else []


def _pptx_native(file_bytes: bytes) -> List[Tuple[int, str]]:
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(file_bytes))
    result: List[Tuple[int, str]] = []
    for slide_index, slide in enumerate(presentation.slides, start=1):
        parts: List[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                text = (shape.text or "").strip()
                if text:
                    is_title = bool(getattr(shape, "is_placeholder", False) and getattr(shape, "placeholder_format", None) and getattr(shape.placeholder_format, "type", None) == 1)
                    parts.append(f"## {text}" if is_title else text)
            if getattr(shape, "has_table", False):
                rows = [
                    " | ".join((cell.text or "").strip().replace("\n", " / ") for cell in row.cells)
                    for row in shape.table.rows
                ]
                if rows:
                    parts.append("TABLE\n" + "\n".join(rows))
        if parts:
            result.append((slide_index, "\n\n".join(parts)))
    return result


def _html_native(file_bytes: bytes) -> List[Tuple[int, str]]:
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(file_bytes, "html.parser")
    parts: List[str] = []
    for element in soup.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "pre", "blockquote", "table"]):
        text = element.get_text(" | " if element.name == "table" else " ", strip=True)
        if not text:
            continue
        if re.fullmatch(r"h[1-6]", element.name):
            text = f"{'#' * (int(element.name[1]) + 1)} {text}"
        elif element.name == "table":
            text = f"TABLE\n{text}"
        parts.append(text)
    content = "\n\n".join(parts).strip()
    return [(0, content)] if content else []


def extract_document_pages(file_bytes: bytes, filename: str) -> List[Tuple[int, str]]:
    """Return (page/slide number, text) pairs with headings and tables preserved."""
    name = os.path.basename(filename or "").strip()
    suffix = os.path.splitext(name)[1].lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type {suffix or '(no extension)'}. Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}.")
    if not file_bytes:
        return []
    ocr_enabled = os.getenv("OCR_ENABLED", "true").strip().lower() in {"1", "true", "yes", "on"}
    ocr_lang = os.getenv("OCR_LANG", "eng")

    if suffix == ".pdf":
        return _pdf_pages(file_bytes, ocr_enabled, ocr_lang)
    if suffix == ".docx":
        try:
            parts = _unstructured_pages(file_bytes, name)
            if any(text.strip() for _, text in parts):
                return parts
        except Exception:
            logger.info("Unstructured DOCX parser unavailable; using python-docx native parsing.")
        return _docx_native(file_bytes)
    if suffix == ".pptx":
        try:
            parts = _unstructured_pages(file_bytes, name)
            if any(text.strip() for _, text in parts):
                return parts
        except Exception:
            logger.info("Unstructured PPTX parser unavailable; using python-pptx native parsing.")
        return _pptx_native(file_bytes)
    if suffix in {".html", ".htm"}:
        try:
            parts = _unstructured_pages(file_bytes, name)
            if any(text.strip() for _, text in parts):
                return parts
        except Exception:
            logger.info("Unstructured HTML parser unavailable; using BeautifulSoup.")
        return _html_native(file_bytes)
    if suffix in {".txt", ".md", ".csv", ".json"}:
        text = file_bytes.decode("utf-8-sig", errors="replace").replace("\x00", "").strip()
        return [(0, text)] if text else []
    return []
